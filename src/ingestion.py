import csv
import json
import re
import shutil
from collections.abc import Iterable
from pathlib import Path
from typing import Protocol

from docx import Document as WordDocument
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation


SUPPORTED_EXTENSIONS = (
    ".pdf",
    ".txt",
    ".docx",
    ".pptx",
    ".xlsx",
    ".csv",
    ".json",
)


class UploadedFile(Protocol):
    name: str

    def getvalue(self) -> bytes: ...


class Ingest:
    def process_uploaded_files(
        self,
        uploaded_files: Iterable[UploadedFile],
        uploads_dir: Path,
    ) -> list[Document]:
        shutil.rmtree(uploads_dir, ignore_errors=True)
        uploads_dir.mkdir(parents=True)

        documents = []
        for uploaded_file in uploaded_files:
            file_name = Path(uploaded_file.name).name
            file_path = uploads_dir / file_name
            file_path.write_bytes(uploaded_file.getvalue())
            documents.extend(self._load_file(file_path))

        return documents

    def process_doc(self, doc_directory: str | Path) -> list[Document]:
        doc_dir = Path(doc_directory)
        documents = []
        for file_path in sorted(doc_dir.rglob("*")):
            if (
                file_path.is_file()
                and file_path.suffix.lower() in SUPPORTED_EXTENSIONS
            ):
                documents.extend(self._load_file(file_path))

        return documents

    @staticmethod
    def _load_file(file_path: Path) -> list[Document]:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            documents = PyPDFLoader(str(file_path)).load()
        elif suffix == ".txt":
            documents = TextLoader(
                str(file_path),
                autodetect_encoding=True,
            ).load()
        elif suffix == ".docx":
            documents = Ingest._load_word_document(file_path)
        elif suffix == ".pptx":
            documents = Ingest._load_powerpoint(file_path)
        elif suffix == ".xlsx":
            documents = Ingest._load_excel(file_path)
        elif suffix == ".csv":
            documents = Ingest._load_csv(file_path)
        elif suffix == ".json":
            documents = Ingest._load_json(file_path)
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

        for document in documents:
            document.metadata["source_file"] = file_path.name
            document.metadata["file_type"] = suffix.lstrip(".")

        return documents

    @staticmethod
    def _load_word_document(file_path: Path) -> list[Document]:
        word_document = WordDocument(str(file_path))
        content = []
        for block in word_document.iter_inner_content():
            if hasattr(block, "text"):
                text = block.text.strip()
                if text:
                    content.append(text)
            else:
                for row in block.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    content.append(" | ".join(cells))

        return [Document(page_content="\n".join(content))]

    @staticmethod
    def _load_powerpoint(file_path: Path) -> list[Document]:
        presentation = Presentation(str(file_path))
        documents = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            content = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    content.append(shape.text.strip())
                if shape.has_table:
                    for row in shape.table.rows:
                        content.append(
                            " | ".join(cell.text.strip() for cell in row.cells)
                        )

            documents.append(
                Document(
                    page_content="\n".join(content),
                    metadata={"page": slide_number},
                )
            )

        return documents

    @staticmethod
    def _load_excel(file_path: Path) -> list[Document]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        documents = []
        try:
            for worksheet in workbook.worksheets:
                rows = list(worksheet.iter_rows(values_only=True))
                if not rows:
                    continue

                populated_counts = [
                    sum(
                        value is not None and bool(str(value).strip())
                        for value in row
                    )
                    for row in rows
                ]
                max_populated = max(populated_counts, default=0)
                header_index = next(
                    (
                        index
                        for index, count in enumerate(populated_counts)
                        if count == max_populated and count >= 2
                    ),
                    None,
                )
                if header_index is None:
                    continue

                headers = [
                    re.sub(r"\s+", " ", str(value)).strip()
                    if value is not None and str(value).strip()
                    else f"Column {index + 1}"
                    for index, value in enumerate(rows[header_index])
                ]

                for row_number, row in enumerate(
                    rows[header_index + 1:],
                    start=header_index + 2,
                ):
                    fields = [
                        f"{headers[index]}: {value}"
                        for index, value in enumerate(row)
                        if (
                            index < len(headers)
                            and value is not None
                            and str(value).strip()
                        )
                    ]
                    if not fields:
                        continue

                    documents.append(
                        Document(
                            page_content="Spreadsheet record: "
                            + " | ".join(fields),
                            metadata={
                                "sheet": worksheet.title,
                                "excel_row": row_number,
                            },
                        )
                    )
        finally:
            workbook.close()

        return documents

    @staticmethod
    def _load_csv(file_path: Path) -> list[Document]:
        with file_path.open(encoding="utf-8-sig", newline="") as csv_file:
            rows = [
                " | ".join(row)
                for row in csv.reader(csv_file)
            ]

        return [Document(page_content="\n".join(rows))]

    @staticmethod
    def _load_json(file_path: Path) -> list[Document]:
        with file_path.open(encoding="utf-8") as json_file:
            data = json.load(json_file)

        if isinstance(data, list):
            records = data
        else:
            records = [data]

        if records and all(isinstance(record, dict) for record in records):
            documents = []
            for record in records:
                fields = [
                    f"{key}: {value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)}"
                    for key, value in record.items()
                ]
                documents.append(
                    Document(
                        page_content="JSON record: " + " | ".join(fields),
                    )
                )

            return documents

        return [
            Document(page_content=json.dumps(data, ensure_ascii=False, indent=2))
        ]
